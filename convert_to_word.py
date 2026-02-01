#!/usr/bin/env python3
"""
Convierte ANEXO_URBANISMO_PROXY.md a formato Word (.docx)
"""

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
import re

def create_word_document(md_file, docx_file):
    """Convierte markdown a Word con formato profesional"""

    # Crear documento
    doc = Document()

    # Configurar estilos
    styles = doc.styles

    # Leer contenido markdown
    with open(md_file, 'r', encoding='utf-8') as f:
        content = f.read()

    lines = content.split('\n')

    in_code_block = False
    in_table = False
    table_lines = []

    i = 0
    while i < len(lines):
        line = lines[i]

        # Bloques de código
        if line.startswith('```'):
            in_code_block = not in_code_block
            i += 1
            continue

        if in_code_block:
            # Añadir línea de código
            p = doc.add_paragraph(line, style='List Bullet')
            p_format = p.paragraph_format
            p_format.left_indent = Inches(0.5)
            run = p.runs[0]
            run.font.name = 'Courier New'
            run.font.size = Pt(9)
            i += 1
            continue

        # Título principal (# )
        if line.startswith('# ') and not line.startswith('## '):
            text = line[2:].strip()
            p = doc.add_heading(text, level=1)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if p.runs:
                run = p.runs[0]
                run.font.size = Pt(18)
                run.font.color.rgb = RGBColor(0, 51, 102)
            i += 1
            continue

        # Subtítulos nivel 2 (## )
        if line.startswith('## '):
            text = line[3:].strip()
            p = doc.add_heading(text, level=2)
            if p.runs:
                run = p.runs[0]
                run.font.size = Pt(14)
                run.font.color.rgb = RGBColor(0, 102, 204)
            i += 1
            continue

        # Subtítulos nivel 3 (### )
        if line.startswith('### '):
            text = line[4:].strip()
            p = doc.add_heading(text, level=3)
            if p.runs:
                run = p.runs[0]
                run.font.size = Pt(12)
            i += 1
            continue

        # Subtítulos nivel 4 (#### )
        if line.startswith('#### '):
            text = line[5:].strip()
            p = doc.add_heading(text, level=4)
            if p.runs:
                run = p.runs[0]
                run.font.size = Pt(11)
            i += 1
            continue

        # Tablas (detectar inicio)
        if '|' in line and not in_table:
            in_table = True
            table_lines = [line]
            i += 1
            continue

        if in_table:
            if '|' in line:
                table_lines.append(line)
                i += 1
                continue
            else:
                # Fin de tabla, procesarla
                add_table_to_doc(doc, table_lines)
                table_lines = []
                in_table = False
                continue

        # Líneas horizontales
        if line.strip() == '---' or line.strip() == '***':
            doc.add_paragraph()
            p = doc.add_paragraph('_' * 80)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            doc.add_paragraph()
            i += 1
            continue

        # Listas con bullets (- o •)
        if line.startswith('- ') or line.startswith('• ') or line.startswith('✓ ') or line.startswith('❌ ') or line.startswith('⚠️ ') or line.startswith('❶ ') or line.startswith('❷ ') or line.startswith('❸ ') or line.startswith('❹ '):
            text = line[2:].strip()
            # Mantener emojis si están presentes
            if line[0] in ['✓', '❌', '⚠', '❶', '❷', '❸', '❹']:
                text = line[0] + ' ' + text
            p = doc.add_paragraph(text, style='List Bullet')
            format_inline_markdown(p)
            i += 1
            continue

        # Listas numeradas
        if re.match(r'^\d+\. ', line):
            text = re.sub(r'^\d+\. ', '', line)
            p = doc.add_paragraph(text, style='List Number')
            format_inline_markdown(p)
            i += 1
            continue

        # Citas/bloques destacados (líneas que empiezan con >)
        if line.startswith('> '):
            text = line[2:].strip()
            p = doc.add_paragraph(text)
            p_format = p.paragraph_format
            p_format.left_indent = Inches(0.5)
            run = p.runs[0]
            run.font.italic = True
            run.font.color.rgb = RGBColor(100, 100, 100)
            i += 1
            continue

        # Párrafos normales
        if line.strip():
            p = doc.add_paragraph(line)
            format_inline_markdown(p)
        else:
            # Línea vacía
            if i > 0 and lines[i-1].strip():  # Solo añadir espacio si no es múltiple
                doc.add_paragraph()

        i += 1

    # Guardar documento
    doc.save(docx_file)
    print(f"✓ Documento Word creado: {docx_file}")


def add_table_to_doc(doc, table_lines):
    """Añade una tabla al documento Word"""

    if len(table_lines) < 2:
        return

    # Filtrar líneas de separación (---)
    data_lines = [l for l in table_lines if not re.match(r'^\|[\s\-:]+\|', l)]

    if len(data_lines) < 1:
        return

    # Parsear filas
    rows = []
    for line in data_lines:
        cells = [c.strip() for c in line.split('|')[1:-1]]  # Ignorar primer y último vacío
        rows.append(cells)

    if not rows:
        return

    # Crear tabla
    num_cols = len(rows[0])
    num_rows = len(rows)

    table = doc.add_table(rows=num_rows, cols=num_cols)
    table.style = 'Light Grid Accent 1'

    # Llenar celdas
    for i, row_data in enumerate(rows):
        row = table.rows[i]
        for j, cell_text in enumerate(row_data):
            cell = row.cells[j]
            cell.text = cell_text

            # Encabezado en negrita
            if i == 0:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.bold = True
                        run.font.size = Pt(10)
            else:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.size = Pt(9)

    doc.add_paragraph()


def format_inline_markdown(paragraph):
    """Formatea markdown inline (**, *, `, etc) en un párrafo"""

    text = paragraph.text
    paragraph.clear()

    # Procesar negritas (**text**)
    parts = re.split(r'(\*\*[^*]+\*\*)', text)
    for part in parts:
        if part.startswith('**') and part.endswith('**'):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        else:
            # Procesar cursivas (*text*)
            sub_parts = re.split(r'(\*[^*]+\*)', part)
            for sub_part in sub_parts:
                if sub_part.startswith('*') and sub_part.endswith('*') and not sub_part.startswith('**'):
                    run = paragraph.add_run(sub_part[1:-1])
                    run.italic = True
                else:
                    # Procesar código (`text`)
                    code_parts = re.split(r'(`[^`]+`)', sub_part)
                    for code_part in code_parts:
                        if code_part.startswith('`') and code_part.endswith('`'):
                            run = paragraph.add_run(code_part[1:-1])
                            run.font.name = 'Courier New'
                            run.font.size = Pt(9)
                            run.font.color.rgb = RGBColor(200, 0, 0)
                        else:
                            paragraph.add_run(code_part)


if __name__ == "__main__":
    md_file = "/home/user/CPQ/ANEXO_URBANISMO_PROXY.md"
    docx_file = "/home/user/CPQ/ANEXO_URBANISMO_PROXY.docx"

    create_word_document(md_file, docx_file)

    print(f"\n✅ Conversión completada")
    print(f"   Archivo: {docx_file}")
