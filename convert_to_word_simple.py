#!/usr/bin/env python3
"""
Convierte ANEXO_URBANISMO_PROXY.md a formato Word (.docx)
Versión simplificada y robusta
"""

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import re

def create_word_document(md_file, docx_file):
    """Convierte markdown a Word con formato profesional"""

    # Crear documento
    doc = Document()

    # Leer contenido markdown
    with open(md_file, 'r', encoding='utf-8') as f:
        content = f.read()

    lines = content.split('\n')

    in_code_block = False
    code_lines = []

    for i, line in enumerate(lines):

        # Bloques de código
        if line.startswith('```'):
            if in_code_block:
                # Terminar bloque de código
                if code_lines:
                    p = doc.add_paragraph()
                    p.style = 'Intense Quote'
                    full_code = '\n'.join(code_lines)
                    run = p.add_run(full_code)
                    run.font.name = 'Courier New'
                    run.font.size = Pt(9)
                    code_lines = []
                in_code_block = False
            else:
                # Iniciar bloque de código
                in_code_block = True
                code_lines = []
            continue

        if in_code_block:
            code_lines.append(line)
            continue

        # Título principal (# )
        if line.startswith('# ') and not line.startswith('## '):
            text = line[2:].strip()
            doc.add_heading(text, level=1)
            continue

        # Subtítulos nivel 2 (## )
        if line.startswith('## '):
            text = line[3:].strip()
            doc.add_heading(text, level=2)
            continue

        # Subtítulos nivel 3 (### )
        if line.startswith('### '):
            text = line[4:].strip()
            doc.add_heading(text, level=3)
            continue

        # Subtítulos nivel 4 (#### )
        if line.startswith('#### '):
            text = line[5:].strip()
            doc.add_heading(text, level=4)
            continue

        # Líneas horizontales
        if line.strip() in ['---', '***', '___']:
            doc.add_paragraph()
            p = doc.add_paragraph('─' * 70)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            doc.add_paragraph()
            continue

        # Tablas
        if '|' in line and '---' not in line:
            # Intentar procesar como parte de tabla
            # (simplificado - solo añadir como texto)
            p = doc.add_paragraph(line)
            run = p.runs[0] if p.runs else p.add_run(line)
            run.font.name = 'Courier New'
            run.font.size = Pt(9)
            continue

        # Listas con bullets
        if line.startswith('- ') or line.startswith('• ') or line.startswith('✓ ') or line.startswith('❌ ') or line.startswith('⚠️ '):
            text = line[2:].strip()
            # Mantener emoji si existe
            if line[0] in ['✓', '❌', '⚠']:
                text = line[0] + ' ' + text
            doc.add_paragraph(text, style='List Bullet')
            continue

        # Símbolos especiales de lista
        if re.match(r'^[❶❷❸❹❺] ', line):
            text = line
            doc.add_paragraph(text, style='List Bullet')
            continue

        # Listas numeradas
        if re.match(r'^\d+\. ', line):
            text = re.sub(r'^\d+\. ', '', line)
            doc.add_paragraph(text, style='List Number')
            continue

        # Párrafos normales
        if line.strip():
            # Procesar texto con formato especial
            if '**' in line or '*' in line or '`' in line:
                p = doc.add_paragraph()
                add_formatted_text(p, line)
            else:
                doc.add_paragraph(line)
        else:
            # Línea vacía - espaciado
            pass

    # Guardar documento
    doc.save(docx_file)
    print(f"✓ Documento Word creado: {docx_file}")


def add_formatted_text(paragraph, text):
    """Añade texto con formato inline (negritas, cursivas, código)"""

    # Procesar negritas (**text**)
    parts = re.split(r'(\*\*[^*]+\*\*)', text)

    for part in parts:
        if part.startswith('**') and part.endswith('**'):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith('`') and part.endswith('`'):
            run = paragraph.add_run(part[1:-1])
            run.font.name = 'Courier New'
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(200, 0, 0)
        else:
            # Procesar cursivas simples
            italic_parts = re.split(r'(\*[^*]+\*)', part)
            for ipart in italic_parts:
                if ipart.startswith('*') and ipart.endswith('*') and len(ipart) > 2:
                    run = paragraph.add_run(ipart[1:-1])
                    run.italic = True
                else:
                    paragraph.add_run(ipart)


if __name__ == "__main__":
    md_file = "/home/user/CPQ/ANEXO_URBANISMO_PROXY.md"
    docx_file = "/home/user/CPQ/ANEXO_URBANISMO_PROXY.docx"

    print("Convirtiendo a Word...")
    create_word_document(md_file, docx_file)

    print(f"\n✅ Conversión completada")
    print(f"   Archivo: {docx_file}")
